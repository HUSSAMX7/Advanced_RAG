import asyncio
import io

import pytest
from docx import Document

from advanced_rag import Settings
from advanced_rag.ingestion.documents import prepare_document


def test_pdf_sections_use_saved_openai_key_and_chat_model(monkeypatch):
    from types import SimpleNamespace

    from llama_index.core.schema import TextNode
    from pydantic import SecretStr

    from advanced_rag.ingestion.chunking import create_chunks

    calls = []

    class FakeOpenAI:
        def __init__(self, *, api_key, **kwargs):
            self.api_key = api_key
            self.responses = self

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def parse(self, *, model, input, text_format):
            calls.append((model, self.api_key))
            return SimpleNamespace(output_parsed=text_format(sections=[]))

    monkeypatch.setattr("advanced_rag.ingestion.sections.AsyncOpenAI", FakeOpenAI)
    settings = Settings(_env_file=None, agent_model="gpt-4o", openai_api_key=SecretStr("web-key"))
    chunks = asyncio.run(
        create_chunks(
            {"notes.pdf": [TextNode(text="Invoice total 120 dollars", metadata={"page_num": 1})]},
            settings=settings,
        )
    )
    assert chunks[0].text == "Invoice total 120 dollars"
    assert calls == [("gpt-4o", "web-key")]


def test_word_text_and_tables_are_chunked_in_document_order_with_file_identity():
    document = Document()
    document.add_paragraph("مقدمة عن المشروع")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ZX-774"
    table.cell(0, 1).text = "120 dollars"
    document.add_paragraph("نهاية المستند")
    output = io.BytesIO()
    document.save(output)
    stages = []
    chunks = asyncio.run(
        prepare_document(
            {"resource_id": "word-id", "file_name": "report.docx", "data": output.getvalue()},
            Settings(_env_file=None),
            stages.append,
        )
    )
    assert stages == ["extracting", "chunking"]
    assert "مقدمة عن المشروع\n\nZX-774 | 120 dollars\n\nنهاية المستند" in chunks[0].text
    assert chunks[0].metadata["resource_id"] == "word-id"
    assert chunks[0].metadata["file_type"] == "docx"
    assert "page_num" not in chunks[0].metadata


@pytest.mark.parametrize("data", [b"\xff", b"   \n "])
def test_invalid_encoding_or_empty_text_fails_before_any_index_write(data):
    with pytest.raises(ValueError):
        asyncio.run(
            prepare_document(
                {"resource_id": "txt", "file_name": "report.txt", "data": data},
                Settings(_env_file=None),
                lambda stage: None,
            )
        )
